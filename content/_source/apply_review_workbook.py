"""Read the reviewer's corrections back out of the review workbook.

Reads column E ("PERBAIKAN / YOUR REVISION") of every sheet, matches each row
to its content by the stable key in column A, and writes the new Indonesian
into the content JSON. English is never touched.

Run a dry run first — it reports what would change without writing anything:

    python3 content/_source/apply_review_workbook.py path/to/reviewed.xlsx
    python3 content/_source/apply_review_workbook.py path/to/reviewed.xlsx --write

Keys look like:
    daily:212:teaching          fw:275:2:looksLike       weekly:7
    guide:9:principle           tool:1:field:3           diary:week:12:theme
    intro:welcome:solo          ui:src/screens/More.tsx:56

Two kinds of key are handled specially:

  daily / fw / weekly / monthly — these live in the generated daily.json, so
  writing there would be undone by the next build. They go into
  content/_source/id_overrides.json, which build_daily.py applies last.

  ui — these live inline in the React components rather than in JSON, so they
  are printed for hand-editing rather than applied automatically.
"""

import json
import pathlib
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(".")
CONTENT = ROOT / "content"
OVERRIDES = ROOT / "content/_source/id_overrides.json"

# Keys whose content is generated into daily.json; these route to the overrides
# file so a rebuild does not discard them.
GENERATED = ("daily", "fw", "weekly", "monthly")


def load(name):
    return json.loads((CONTENT / name).read_text(encoding="utf-8"))


def save(name, data):
    (CONTENT / name).write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def resolve(key, docs):
    """Return (container, field) so that container[field]['id'] is the target."""
    p = key.split(":")
    kind = p[0]

    if kind == "daily":
        day = next(x for x in docs["daily"]["days"] if x["day"] == int(p[1]))
        return day, p[2]

    if kind == "fw":
        day = next(x for x in docs["daily"]["days"] if x["day"] == int(p[1]))
        fw = day["framework"]
        if p[2] in ("title", "note"):
            return fw, p[2]
        entry = fw["entries"][int(p[2])]
        if p[3] == "tag":
            return entry["tags"], int(p[4])
        return entry, p[3]

    if kind in ("weekly", "monthly"):
        bucket = "weeklyExercises" if kind == "weekly" else "monthlyReviews"
        item = next(x for x in docs["daily"][bucket] if x["n"] == int(p[1]))
        return item, "text"

    if kind == "guide":
        ch = next(x for x in docs["guidebook"]["chapters"] if x["number"] == int(p[1]))
        if p[2] == "dailyPractices":
            return ch["dailyPractices"]["id"], int(p[3])  # plain list of strings
        return ch, p[2]

    if kind == "tool":
        if p[1] == "meta":
            return docs["toolkit"]["meta"], p[2]
        tool = next(x for x in docs["toolkit"]["tools"] if x["number"] == int(p[1]))
        if p[2] == "guide":
            return tool["guide"]["id"], int(p[3])       # plain list of strings
        if p[2] == "field":
            return tool["fields"][int(p[3])], "label"
        if p[2] == "summaryField":
            return tool["summaryFields"][int(p[3])], "label"
        if p[2] == "repeat":
            return tool["repeat"], p[3]
        return tool, p[2]

    if kind == "diary":
        m = docs["diary_weeks"]["meta"]
        if p[1] == "dailyPrompt":
            return m["dailyPrompts"], int(p[2])
        if p[1] == "debrief":
            return m["debrief"][p[2]], int(p[3])
        wk = next(x for x in docs["diary_weeks"]["weeks"] if x["week"] == int(p[2]))
        if p[3] == "sunday":
            return wk["sundayReflection"], p[4]
        return wk, p[3]

    if kind == "intro":
        i = docs["diary_intro"]
        if p[1] == "welcome":
            return i["welcome"], p[2]
        if p[1] == "howToUse":
            return i["howToUse"][p[2]][int(p[3])], p[4]
        if p[2] == "intro":
            return i["settingUp"], "intro"
        return i["settingUp"][p[2]]["prompts"], int(p[3])

    raise KeyError(key)


def current_english(docs):
    """key -> the English the app currently shows, for the stale-file guard."""
    import build_review_workbook as base

    return {row[0]: row[2] for _t, _s, rows, _n, _p in base.collect() for row in rows}


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    path = sys.argv[1]
    write = "--write" in sys.argv

    docs = {
        "daily": load("daily.json"),
        "guidebook": load("guidebook.json"),
        "toolkit": load("toolkit.json"),
        "diary_weeks": load("diary_weeks.json"),
        "diary_intro": load("diary_intro.json"),
    }
    files = {
        "daily": "daily.json", "fw": "daily.json", "weekly": "daily.json",
        "monthly": "daily.json", "guide": "guidebook.json", "tool": "toolkit.json",
        "diary": "diary_weeks.json", "intro": "diary_intro.json",
    }

    overrides = json.loads(OVERRIDES.read_text(encoding="utf-8")) if OVERRIDES.exists() else {}

    live_en = current_english(docs)

    wb = load_workbook(path, read_only=True, data_only=True)
    applied, ui, unchanged, failed, stale = [], [], 0, [], []
    touched = set()
    n_overrides = 0

    for ws in wb.worksheets:
        if ws.title == "Baca dulu":
            continue
        for row in ws.iter_rows(min_row=3, max_col=5, values_only=True):
            key, _where, en, current, revision = row
            if not key or not revision or not str(revision).strip():
                continue
            revision = str(revision).strip()
            if revision == (current or "").strip():
                unchanged += 1
                continue
            # The English in the workbook must still be the English in the app.
            # If it is not, this row was reviewed against content that has since
            # been rewritten, and its Indonesian would land on a different
            # sentence — so refuse it rather than create a mismatch.
            live = live_en.get(key)
            if live is not None and en is not None and live.strip() != str(en).strip():
                stale.append((key, str(en).strip(), live.strip()))
                continue
            if key.startswith("ui:"):
                ui.append((key, current, revision))
                continue
            if key.split(":")[0] in GENERATED:
                overrides[key] = revision
                n_overrides += 1
                continue
            try:
                container, field = resolve(key, docs)
            except (KeyError, StopIteration, IndexError):
                failed.append(key)
                continue
            target = container[field]
            if isinstance(target, dict):        # {'en': …, 'id': …}
                target["id"] = revision
            else:                               # plain string inside an id list
                container[field] = revision
            applied.append(key)
            touched.add(files[key.split(":")[0]])

    total = len(applied) + n_overrides + len(ui) + unchanged + len(failed) + len(stale)
    print(f"revisions found      : {total}")
    print(f"  applied to JSON    : {len(applied)}")
    print(f"  daily overrides    : {n_overrides}")
    print(f"  UI (edit by hand)  : {len(ui)}")
    print(f"  identical, skipped : {unchanged}")
    print(f"  unrecognised keys  : {len(failed)}")
    print(f"  STALE, not applied : {len(stale)}")
    if stale:
        print("\n  These rows were reviewed against English that has since changed.")
        print("  Applying them would put the Indonesian on a different sentence.")
        print("  Rebuild the workbook and have those lines reviewed again.\n")
        for key, was, now in stale[:5]:
            print(f"    {key}\n      workbook English: {was[:80]}\n      current English : {now[:80]}")
        if len(stale) > 5:
            print(f"    … and {len(stale) - 5} more")
    for k in failed[:10]:
        print(f"      {k}")
    # A UI wording can appear in more than one component. The workbook shows it
    # once, so name every location that still carries the old text — editing
    # only the keyed one leaves the same English rendering two ways.
    strings = json.loads(
        (ROOT / "content/_source/ui_strings.json").read_text(encoding="utf-8")
    ) if (ROOT / "content/_source/ui_strings.json").exists() else []
    for key, before, after in ui:
        _, path, line = key.split(":")[0], ":".join(key.split(":")[1:-1]), key.split(":")[-1]
        also = [
            f"{s['file']}:{s['line']}"
            for s in strings
            if s["id"] == before and f"{s['file']}:{s['line']}" != f"{path}:{line}"
        ]
        print(f"\n  {path}:{line}\n    - {before}\n    + {after}")
        for loc in also:
            print(f"    ALSO HERE: {loc}")

    if not write:
        print("\nDry run — nothing written. Re-run with --write to apply.")
        return
    for name in sorted(touched):
        stem = name.replace(".json", "")
        save(name, docs[stem])
        print(f"wrote content/{name}")
    if n_overrides:
        OVERRIDES.write_text(
            json.dumps(overrides, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        print(f"wrote {OVERRIDES}  ({len(overrides)} total)")
        print("Now run: python3 content/_source/build_daily.py")


if __name__ == "__main__":
    main()
